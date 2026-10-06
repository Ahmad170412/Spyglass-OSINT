import os
import json
import mimetypes
import subprocess
import hashlib
from datetime import datetime

from .utils import _run, _check_tool, _EXIFTOOL


def normalise_path(raw):
    """Turn a pasted or typed path into one this OS can open.

    Three front ends hand paths to this module — the one-shot CLI, the
    interactive menu (which has no shell behind it, so ``~`` reaches us as a
    literal tilde), and the web console, where paths arrive from a text box
    carrying whatever the file manager put on the clipboard, quotes included.
    Doing it once here means none of them has to know this platform's separator,
    home-directory syntax or variable form.
    """
    p = str(raw or "").strip()
    # Finder, Explorer and most file managers quote a path containing spaces.
    if len(p) >= 2 and p[0] == p[-1] and p[0] in "\"'":
        p = p[1:-1].strip()
    # expandvars reads %VAR% on Windows and $VAR elsewhere; expanduser reads
    # the platform's home syntax. Both are no-ops when there is nothing to expand.
    return os.path.expandvars(os.path.expanduser(p))


def extract(path):
    """Extract metadata from a file. Returns a structured dict."""
    path = normalise_path(path)
    if not path:
        return {"error": "No file path given"}
    if not os.path.isfile(path):
        return {"error": f"File not found: {path}"}

    result = _file_info(path)

    sniffed = None
    if _EXIFTOOL:
        meta, sniffed = _exiftool_extract(path)
        result["file_type"] = _resolve_type(sniffed, result["file_type"])
        if meta:
            result.update(meta)
            return result

    # No exiftool — try Python fallbacks per type
    ext = os.path.splitext(path)[1].lower()
    if ext in (".jpg", ".jpeg", ".tiff", ".tif", ".png", ".webp"):
        img = _pillow_exif(path)
        if img:
            result.update(img)
    elif ext == ".pdf":
        pdf = _pdf_meta(path)
        if pdf:
            result.update(pdf)
    elif ext in (".docx", ".xlsx", ".pptx"):
        ooxml = _ooxml_meta(path)
        if ooxml:
            result.update(ooxml)

    return result


def _file_info(path):
    st = os.stat(path)
    hashes = _compute_hashes(path)
    return {
        "file_name": os.path.basename(path),
        "file_size": _human_size(st.st_size),
        "file_type": _mime_type(path),
        "created": _created(st),
        "modified": datetime.fromtimestamp(st.st_mtime).isoformat(),
        "accessed": datetime.fromtimestamp(st.st_atime).isoformat(),
        "md5": hashes["md5"],
        "sha256": hashes["sha256"],
    }


# Read once rather than per call so the platform branch in `_created` is a
# module constant a test can flip, instead of a patch on the global `os`.
_IS_WINDOWS = os.name == "nt"


def _created(st):
    """Creation time, with an honest answer on every platform this runs on.

    ``st_birthtime`` is macOS/BSD. Linux does not expose it at all, so there is
    no creation time to report and ``N/A`` is the right value rather than a
    silent fallback to the inode-change time, which is not the same fact. On
    Windows ``st_ctime`` *is* the creation time (its change time is ``st_mtime``),
    so it is used there instead of reporting nothing.
    """
    birth = getattr(st, "st_birthtime", None)
    if birth is None and _IS_WINDOWS:
        birth = st.st_ctime
    if birth is None:
        return "N/A"
    return datetime.fromtimestamp(birth).isoformat()


def _compute_hashes(path):
    """Compute MD5 and SHA256 hashes of a file."""
    md5_hash = hashlib.md5()
    sha256_hash = hashlib.sha256()
    try:
        with open(path, "rb") as f:
            for chunk in iter(lambda: f.read(8192), b""):
                md5_hash.update(chunk)
                sha256_hash.update(chunk)
        return {"md5": md5_hash.hexdigest(), "sha256": sha256_hash.hexdigest()}
    except Exception:
        return {"md5": "N/A", "sha256": "N/A"}


def _exiftool_extract(path):
    """Format metadata via exiftool — the gold standard for a format it knows.

    Returns ``(fields, sniffed_mime)``. The MIME type is pulled out separately
    rather than left in the fields: exiftool's own answer for the file's
    *identity* still has to be reconciled with the extension table (see
    ``_resolve_type``), and it must never be reported twice. ``(None, None)``
    means the run failed, as distinct from ``(dict(), ...)`` which means it
    succeeded and simply had nothing curated to show.
    """
    try:
        r = subprocess.run(
            [_EXIFTOOL, "-j", "-G", path],
            capture_output=True, text=True, timeout=30,
        )
        data = json.loads(r.stdout.strip())
        if isinstance(data, list) and data:
            raw = data[0]
            return _clean_exiftool(raw), _exiftool_mime(raw)
    except Exception:
        pass
    return None, None


def _exiftool_mime(raw):
    for key, val in raw.items():
        if key.rsplit(":", 1)[-1] == "MIMEType" and val:
            return str(val).strip()
    return None


_CURATED = {
    # File. Deliberately *not* FileSize / FileType / FileTypeExtension /
    # MIMEType: `_file_info` already reports those four facts, from stat and
    # from the extension, and keeping both copies is what put two contradictory
    # "file type" rows on screen — a .py script came back as TXT from exiftool
    # and application/octet-stream from us. `_resolve_type` folds exiftool's
    # MIME into `file_type` where it adds something.
    "ImageWidth", "ImageHeight", "Megapixels", "Duration",
    # Image
    "Make", "Model", "Software", "Orientation", "XResolution", "YResolution",
    "ISO", "Aperture", "FNumber", "ExposureTime", "FocalLength",
    "Flash", "WhiteBalance", "ColorSpace",
    # Timestamps
    "DateTimeOriginal", "CreateDate", "ModifyDate", "DateCreated",
    "ProfileDateTime", "CreationDate",
    # GPS
    "GPSLatitude", "GPSLatitudeRef", "GPSLongitude", "GPSLongitudeRef",
    "GPSAltitude", "GPSAltitudeRef", "GPSPosition", "GPSDateTime",
    # Documents
    "Author", "Creator", "Producer", "Title", "Subject",
    "Description", "Keywords", "LastModifiedBy", "RevisionNumber",
    "Company", "Manager", "Category", "Language", "Pages",
    # Audio/Video
    "Artist", "Album", "Title", "Track", "Genre", "Duration",
    "SampleRate", "Channels", "AudioBitrate",
    # Network / forensic
    "UserComment", "Warning",
}

_IGNORE_VALUES = {"N/A", "-", "unknown", "Unspecified", "", "none", "None"}


def _clean_exiftool(raw):
    """Return only curated fields + GPS, with clean names."""
    out = {}
    gps = {}

    for key, val in raw.items():
        if val is None or val == "":
            continue
        s = str(val).strip()
        if s in _IGNORE_VALUES:
            continue
        short = key.split(":")[-1] if ":" in key else key

        # Collect GPS subfields separately
        if short.startswith("GPS") and short != "GPSPosition":
            gps[short] = s
            continue

        if short not in _CURATED:
            continue

        out[short] = s

    if gps:
        out["GPS"] = gps
    return out


def _pillow_exif(path):
    """Image EXIF via Pillow."""
    try:
        from PIL import Image
        from PIL.ExifTags import TAGS
    except ImportError:
        return {}

    try:
        img = Image.open(path)
        exif = img._getexif()
        if not exif:
            return {}
        out = {}
        gps = {}
        for tag_id, val in exif.items():
            name = TAGS.get(tag_id, tag_id)
            if isinstance(val, bytes):
                try:
                    val = val.decode("utf-8", errors="replace").strip()
                except Exception:
                    continue
            if name == "GPSInfo":
                for k, v in val.items():
                    gps_tag = {0: "GPSLatitudeRef", 1: "GPSLatitude",
                               2: "GPSLongitudeRef", 3: "GPSLongitude",
                               4: "GPSAltitudeRef", 5: "GPSAltitude"}.get(k, str(k))
                    gps[gps_tag] = v
                continue
            if isinstance(val, (int, float)):
                out[name] = val
            else:
                s = str(val).strip()
                if s and s != "N/A":
                    out[name] = s
        if gps:
            out["GPS"] = gps
        return out
    except Exception:
        return {}


def _pdf_meta(path):
    """PDF metadata via PyPDF2."""
    try:
        import PyPDF2
    except ImportError:
        try:
            import pdfminer
            return _pdfminer_meta(path)
        except ImportError:
            return {}
    try:
        with open(path, "rb") as f:
            r = PyPDF2.PdfReader(f)
            info = r.metadata
            if info:
                return {k: str(v) for k, v in info.items() if v}
    except Exception:
        pass
    return {}


def _pdfminer_meta(path):
    try:
        from pdfminer.pdfparser import PDFParser
        from pdfminer.pdfdocument import PDFDocument
        with open(path, "rb") as f:
            parser = PDFParser(f)
            doc = PDFDocument(parser)
            return {k: str(v) for k, v in (doc.info[0] if doc.info else {}).items() if v}
    except Exception:
        return {}


def _ooxml_meta(path):
    """Office Open XML metadata (docx/xlsx/pptx) via python-docx/openpyxl."""
    ext = os.path.splitext(path)[1].lower()
    try:
        if ext == ".docx":
            import docx
            d = docx.Document(path)
            p = d.core_properties
        elif ext == ".xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(path)
            p = wb.properties
        else:
            return {}
        out = {}
        for attr in ("author", "category", "comments", "content_status", "created",
                      "creator", "description", "identifier", "keywords", "language",
                      "last_modified_by", "last_printed", "modified", "revision",
                      "subject", "title", "version"):
            val = getattr(p, attr, None)
            if val:
                out[attr.replace("_", " ").title()] = str(val)
        return out
    except Exception:
        return {}


def _human_size(b):
    for unit in ("B", "KB", "MB", "GB"):
        if b < 1024:
            return f"{b:.1f} {unit}"
        b /= 1024
    return f"{b:.1f} TB"


# Extension → MIME, consulted *before* the stdlib because `mimetypes` is filled
# from the host OS — its registry on Windows, /etc/mime.types on Linux — so the
# same file is typed differently depending on where Spyglass runs, and `.rs`
# comes back as `application/rls-services+xml` there, which is an XML format
# rather than Rust. Anything not named here defers to it, then to octet-stream.
_MIME_BY_EXT = {
    # images
    ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png",
    ".gif": "image/gif", ".webp": "image/webp", ".tiff": "image/tiff",
    ".tif": "image/tiff", ".bmp": "image/bmp", ".svg": "image/svg+xml",
    # documents
    ".pdf": "application/pdf",
    ".doc": "application/msword",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".xls": "application/vnd.ms-excel",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".ppt": "application/vnd.ms-powerpoint",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    # audio / video
    ".mp3": "audio/mpeg", ".wav": "audio/wav", ".m4a": "audio/mp4",
    ".ogg": "audio/ogg", ".flac": "audio/flac",
    ".mp4": "video/mp4", ".mov": "video/quicktime", ".mkv": "video/x-matroska",
    # archives
    ".zip": "application/zip", ".gz": "application/gzip",
    ".tar": "application/x-tar", ".7z": "application/x-7z-compressed",
    # source. The extensions exiftool most often has no format entry for, which
    # is exactly how a Python script ends up typed TXT.
    ".py": "text/x-python", ".pyw": "text/x-python", ".pyi": "text/x-python",
    ".rb": "text/x-ruby", ".pl": "text/x-perl", ".php": "text/x-php",
    ".js": "text/javascript", ".mjs": "text/javascript",
    ".cjs": "text/javascript", ".ts": "text/x-typescript",
    ".sh": "application/x-sh", ".bash": "application/x-sh",
    ".zsh": "application/x-sh", ".fish": "application/x-sh",
    ".c": "text/x-c", ".h": "text/x-c", ".cc": "text/x-c++",
    ".cpp": "text/x-c++", ".hpp": "text/x-c++", ".cs": "text/x-csharp",
    ".java": "text/x-java", ".kt": "text/x-kotlin", ".go": "text/x-go",
    ".rs": "text/x-rust", ".swift": "text/x-swift", ".scala": "text/x-scala",
    ".r": "text/x-r", ".lua": "text/x-lua", ".ps1": "text/x-powershell",
    ".sql": "application/sql",
    # data and markup
    ".json": "application/json", ".xml": "application/xml",
    ".yaml": "application/yaml", ".yml": "application/yaml",
    ".toml": "application/toml", ".ini": "text/plain",
    ".cfg": "text/plain", ".conf": "text/plain",
    ".md": "text/markdown", ".rst": "text/x-rst", ".csv": "text/csv",
    ".tsv": "text/tab-separated-values", ".log": "text/plain",
    ".txt": "text/plain", ".html": "text/html", ".htm": "text/html",
    ".css": "text/css",
}


def _mime_type(path):
    """The file's MIME type, from the extension. Same answer on every OS."""
    ext = os.path.splitext(path)[1].lower()
    if ext in _MIME_BY_EXT:
        return _MIME_BY_EXT[ext]
    return mimetypes.guess_type(path)[0] or "application/octet-stream"


# What exiftool falls back to when it has no format registered for the
# extension. It describes the *bytes*, not the file: a Python script, a README
# and a shell script all sniff as text/plain.
_GENERIC_MIMES = {"application/octet-stream", "text/plain"}


def _resolve_type(sniffed, ours):
    """One answer for "what kind of file is this".

    Exiftool is exact for a format it has registered and blind for one it has
    not: it knows `.sh` but not `.py`, so a Python script is sniffed as plain
    text and reported TXT / txt / text/plain. The extension table is the
    reverse — exact for the extensions it names, octet-stream for the rest.
    Neither wins outright, so each is trusted only where it is informative:
    exiftool's MIME when it recognised a real format, or when ours is the empty
    answer; ours whenever ours names something specific.
    """
    if not sniffed or sniffed == ours:
        return ours
    if sniffed not in _GENERIC_MIMES:
        return sniffed
    if ours in _GENERIC_MIMES:
        return sniffed
    return ours
